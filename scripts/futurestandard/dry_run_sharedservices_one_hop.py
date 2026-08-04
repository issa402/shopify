"""
USE CASE: Dry run from SharedServices or another account that can directly assume member roles.

Run this when your terminal credentials can already assume each member account's
ManagedInstance-CrossAccountExecutionRole. It only checks key status and writes
a report. It does not disable anything.
"""

# =========================
# PART 1: IMPORTS
# =========================
import csv  # Writes the dry-run report CSV.
from collections import defaultdict  # Groups workbook rows by AWS account.
from datetime import datetime  # Adds timestamps to result rows.

import boto3  # AWS SDK for Python.
from openpyxl import load_workbook  # Reads the Excel workbook.


# =========================
# PART 2: SETTINGS
# =========================
EXCEL_FILE = "Unused_Keys_By-Days.xlsx"  # Workbook with account/user/key/reason.
RESULTS_CSV = "dry_run_sharedservices_one_hop_results.csv"  # Output report.
MEMBER_ACCOUNT_ROLE_NAME = "ManagedInstance-CrossAccountExecutionRole"  # Member role.
AWS_ACCOUNT_ID_LENGTH = 12  # AWS account IDs are always 12 digits.


# =========================
# PART 3: CLEAN ACCOUNT IDS
# =========================
def clean_account_id(value):
    account = str(value).strip()  # Convert Excel value to text and trim spaces.
    if account.endswith(".0"):  # Excel may format account IDs as 123.0.
        account = account[:-2]  # Remove .0.
    account = "".join(char for char in account if char.isdigit())  # Keep only digits.
    return account.zfill(AWS_ACCOUNT_ID_LENGTH)  # Restore missing leading zeroes.


# =========================
# PART 4: LOAD WORKBOOK ROWS
# =========================
def load_rows():
    wb = load_workbook(EXCEL_FILE, data_only=True)  # Open workbook values.
    rows = []  # Store cleaned rows.

    for sheet in wb.worksheets:  # Read each sheet.
        headers = {}  # Header name -> column number.
        for col_num, cell in enumerate(sheet[1], start=1):  # Header row.
            headers[str(cell.value).strip()] = col_num  # Save header mapping.

        for row_num in range(2, sheet.max_row + 1):  # Data rows.
            account = sheet.cell(row=row_num, column=headers["account"]).value  # Account ID.
            user_name = sheet.cell(row=row_num, column=headers["user_name"]).value  # IAM user.
            access_key_id = sheet.cell(row=row_num, column=headers["access_key_id"]).value  # Key ID.
            reason = sheet.cell(row=row_num, column=headers["reason"]).value  # Reason.

            if not account or not user_name or not access_key_id:  # Skip incomplete rows.
                continue  # Next row.

            rows.append({
                "sheet": sheet.title,  # Source sheet.
                "row_number": row_num,  # Source row.
                "account": clean_account_id(account),  # 12-digit account.
                "user_name": str(user_name).strip(),  # Clean user.
                "access_key_id": str(access_key_id).strip(),  # Clean key.
                "reason": str(reason).strip() if reason else "",  # Clean reason.
            })

    return rows  # Return all rows.


# =========================
# PART 5: ASSUME MEMBER ROLE
# =========================
def assume_member_account(account_id):
    sts = boto3.client("sts")  # Uses current terminal credentials.
    role_arn = f"arn:aws:iam::{account_id}:role/{MEMBER_ACCOUNT_ROLE_NAME}"  # Member role ARN.
    print(f"  Assuming member account role: {role_arn}")  # Show role.

    response = sts.assume_role(  # Ask STS for member-account credentials.
        RoleArn=role_arn,  # Role to assume.
        RoleSessionName=f"dry-run-unused-keys-{account_id}",  # CloudTrail name.
    )
    creds = response["Credentials"]  # Temporary credentials.

    return boto3.client(  # IAM client inside member account.
        "iam",
        aws_access_key_id=creds["AccessKeyId"],
        aws_secret_access_key=creds["SecretAccessKey"],
        aws_session_token=creds["SessionToken"],
    )


# =========================
# PART 6: DRY-RUN KEY CHECK
# =========================
def get_key_status(iam, user_name, access_key_id):
    response = iam.list_access_keys(UserName=user_name)  # Read keys for user.
    for key in response["AccessKeyMetadata"]:  # Loop through keys.
        if key["AccessKeyId"] == access_key_id:  # Match workbook key.
            return key["Status"]  # Active or Inactive.
    return "NOT_FOUND"  # Key not found.


# =========================
# PART 7: GROUP ROWS
# =========================
def group_rows_by_account(rows):
    grouped = defaultdict(list)  # account_id -> rows.
    for row in rows:  # Each workbook row.
        grouped[row["account"]].append(row)  # Group by account.
    return grouped  # Return grouped rows.


# =========================
# PART 8: MAIN PROGRAM
# =========================
def main():
    rows = load_rows()  # Read workbook.
    grouped_rows = group_rows_by_account(rows)  # Group rows by account.
    results = []  # Store output rows.

    for account_id, account_rows in grouped_rows.items():  # Account by account.
        print(f"\nChecking account {account_id}")  # Progress.
        try:
            iam = assume_member_account(account_id)  # Enter member account.
            for row in account_rows:  # Check keys in account.
                try:
                    key_status = get_key_status(iam, row["user_name"], row["access_key_id"])
                    if key_status == "Active":
                        action, status, error = "DRY_RUN_WOULD_DISABLE", "SUCCESS", ""
                    elif key_status == "Inactive":
                        action, status, error = "SKIP_ALREADY_INACTIVE", "SUCCESS", ""
                    else:
                        action, status, error = "SKIP_KEY_NOT_FOUND", "FAILED", "Access key not found for this user."
                except Exception as e:
                    key_status, action, status, error = "", "FAILED", "FAILED", str(e)

                results.append({**row, "key_status": key_status, "action": action, "status": status, "error": error, "timestamp": datetime.utcnow().isoformat()})
                print(f"{action}: {row['account']} {row['user_name']} {row['access_key_id']}")
        except Exception as e:
            for row in account_rows:
                results.append({**row, "key_status": "", "action": "FAILED_ASSUME_ROLE", "status": "FAILED", "error": str(e), "timestamp": datetime.utcnow().isoformat()})
            print(f"FAILED ACCOUNT {account_id}: {e}")

    with open(RESULTS_CSV, "w", newline="", encoding="utf-8") as f:  # Write CSV.
        writer = csv.DictWriter(f, fieldnames=list(results[0].keys()))  # Headers.
        writer.writeheader()  # Header row.
        writer.writerows(results)  # Result rows.

    print(f"\nChecked {len(results)}")  # Total rows.
    print(f"Results to {RESULTS_CSV}")  # Output file.


if __name__ == "__main__":  # Only run when executed directly.
    main()  # Start script.
