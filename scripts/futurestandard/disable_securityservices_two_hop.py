"""
USE CASE: Real disable run from SecurityServices.

Run this when your terminal credentials are for SecurityServices. The script
assumes SharedServices first, then each member account, then disables the exact
access keys listed in the workbook by setting them to Inactive.
"""

# =========================
# PART 1: IMPORTS
# =========================
import csv  # Writes the final results CSV.
from collections import defaultdict  # Groups workbook rows by AWS account.
from datetime import datetime  # Adds timestamps to each result row.

import boto3  # AWS SDK for Python.
from openpyxl import load_workbook  # Reads the Excel workbook.


# =========================
# PART 2: SETTINGS
# =========================
EXCEL_FILE = "Unused_Keys_By-Days.xlsx"  # Workbook with account/user/key/reason.
RESULTS_CSV = "disabled_securityservices_two_hop_results.csv"  # Output report.

SHARED_SERVICES_ACCOUNT_ID = "123912334908"  # SharedServices account.
SHARED_SERVICES_ROLE_NAME = "ManagedInstance-CrossAccountManagementRole"  # First role.
MEMBER_ACCOUNT_ROLE_NAME = "ManagedInstance-CrossAccountExecutionRole"  # Member role.
AWS_ACCOUNT_ID_LENGTH = 12  # AWS account IDs are always 12 digits.


# =========================
# PART 3: CLEAN ACCOUNT IDS
# =========================
def clean_account_id(value):
    account = str(value).strip()  # Convert Excel account value to text.
    if account.endswith(".0"):  # Excel may format account IDs as 123.0.
        account = account[:-2]  # Remove .0.
    account = "".join(char for char in account if char.isdigit())  # Keep digits only.
    return account.zfill(AWS_ACCOUNT_ID_LENGTH)  # Restore missing leading zeroes.


# =========================
# PART 4: LOAD WORKBOOK ROWS
# =========================
def load_rows():
    wb = load_workbook(EXCEL_FILE, data_only=True)  # Open workbook values.
    rows = []  # Store cleaned rows.

    for sheet in wb.worksheets:  # Read every worksheet.
        headers = {}  # Header name -> column number.
        for col_num, cell in enumerate(sheet[1], start=1):  # Header row.
            headers[str(cell.value).strip()] = col_num  # Save header mapping.

        for row_num in range(2, sheet.max_row + 1):  # Data rows.
            account = sheet.cell(row=row_num, column=headers["account"]).value  # Account ID.
            user_name = sheet.cell(row=row_num, column=headers["user_name"]).value  # IAM user.
            access_key_id = sheet.cell(row=row_num, column=headers["access_key_id"]).value  # Key ID.
            reason = sheet.cell(row=row_num, column=headers["reason"]).value  # Reason.

            if not account or not user_name or not access_key_id:  # Skip blank rows.
                continue  # Next row.

            rows.append({
                "sheet": sheet.title,  # Source sheet.
                "row_number": row_num,  # Source Excel row.
                "account": clean_account_id(account),  # 12-digit account.
                "user_name": str(user_name).strip(),  # Clean IAM user.
                "access_key_id": str(access_key_id).strip(),  # Clean key ID.
                "reason": str(reason).strip() if reason else "",  # Clean reason.
            })

    return rows  # Return workbook rows.


# =========================
# PART 5: AWS ASSUME ROLE HELPERS
# =========================
def assume_role(sts_client, role_arn, session_name):
    response = sts_client.assume_role(  # Ask STS for temporary credentials.
        RoleArn=role_arn,  # Full role ARN.
        RoleSessionName=session_name,  # CloudTrail session name.
    )
    return response["Credentials"]  # Temporary access key/secret/token.


def make_client(service_name, credentials):
    return boto3.client(  # Build a client from temporary credentials.
        service_name,  # Example: sts or iam.
        aws_access_key_id=credentials["AccessKeyId"],  # Temporary access key.
        aws_secret_access_key=credentials["SecretAccessKey"],  # Temporary secret.
        aws_session_token=credentials["SessionToken"],  # Temporary token.
    )


def assume_shared_services():
    sts = boto3.client("sts")  # Uses SecurityServices terminal credentials.
    identity = sts.get_caller_identity()  # Confirm starting identity.
    print(f"Starting identity: {identity['Arn']}")  # Show identity.

    role_arn = f"arn:aws:iam::{SHARED_SERVICES_ACCOUNT_ID}:role/{SHARED_SERVICES_ROLE_NAME}"
    print(f"Assuming SharedServices role: {role_arn}")  # First hop.
    return assume_role(sts, role_arn, "disable-unused-keys-sharedservices")  # SecurityServices -> SharedServices.


def assume_member_account(shared_credentials, account_id):
    shared_sts = make_client("sts", shared_credentials)  # STS client as SharedServices.
    role_arn = f"arn:aws:iam::{account_id}:role/{MEMBER_ACCOUNT_ROLE_NAME}"  # Member role ARN.
    print(f"  Assuming member account role: {role_arn}")  # Second hop.
    member_credentials = assume_role(shared_sts, role_arn, f"disable-unused-keys-{account_id}")
    return make_client("iam", member_credentials)  # IAM client inside member account.


# =========================
# PART 6: DISABLE KEY
# =========================
def disable_key(iam, user_name, access_key_id):
    iam.update_access_key(  # This is the real AWS change.
        UserName=user_name,  # IAM user from workbook.
        AccessKeyId=access_key_id,  # Exact key from workbook.
        Status="Inactive",  # Disable key without deleting it.
    )


# =========================
# PART 7: GROUP ROWS
# =========================
def group_rows_by_account(rows):
    grouped = defaultdict(list)  # account_id -> rows.
    for row in rows:  # Each workbook row.
        grouped[row["account"]].append(row)  # Group by account.
    return grouped  # Return grouped rows.


# =========================
# PART 8: MAIN PROGRAM
# =========================
def main():
    rows = load_rows()  # Read workbook.
    grouped_rows = group_rows_by_account(rows)  # Assume each account once.
    shared_credentials = assume_shared_services()  # First hop.
    results = []  # Store output rows.

    for account_id, account_rows in grouped_rows.items():  # Account by account.
        print(f"\nProcessing account {account_id}")  # Progress.
        try:
            iam = assume_member_account(shared_credentials, account_id)  # Second hop.
            for row in account_rows:  # Disable each key in account.
                try:
                    disable_key(iam, row["user_name"], row["access_key_id"])  # Real disable.
                    action, status, error = "DISABLED", "SUCCESS", ""
                except Exception as e:
                    action, status, error = "FAILED_DISABLE", "FAILED", str(e)

                results.append({**row, "action": action, "status": status, "error": error, "timestamp": datetime.utcnow().isoformat()})
                print(f"{status}: {row['account']} {row['user_name']} {row['access_key_id']}")
        except Exception as e:
            for row in account_rows:
                results.append({**row, "action": "FAILED_ASSUME_ROLE", "status": "FAILED", "error": str(e), "timestamp": datetime.utcnow().isoformat()})
            print(f"FAILED ACCOUNT {account_id}: {e}")

    with open(RESULTS_CSV, "w", newline="", encoding="utf-8") as f:  # Write CSV.
        writer = csv.DictWriter(f, fieldnames=list(results[0].keys()))  # Headers.
        writer.writeheader()  # Header row.
        writer.writerows(results)  # Results.

    print(f"\nProcessed {len(results)}")  # Total rows.
    print(f"Results to {RESULTS_CSV}")  # Output file.


if __name__ == "__main__":  # Only run when executed directly.
    main()  # Start script.
