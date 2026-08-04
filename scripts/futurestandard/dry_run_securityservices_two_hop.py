"""
USE CASE: Dry run from SecurityServices.

Run this when your terminal credentials are for SecurityServices. The script
assumes the SharedServices management role first, then assumes each member
account execution role, then only checks whether each access key exists and is
Active or Inactive. It does not disable anything.
"""

# =========================
# PART 1: IMPORTS
# =========================
import csv  # Writes the dry-run report CSV.
from collections import defaultdict  # Groups workbook rows by AWS account.
from datetime import datetime  # Adds timestamps to each result row.

import boto3  # AWS SDK for Python.
from openpyxl import load_workbook  # Reads the Excel workbook.


# =========================
# PART 2: SETTINGS
# =========================
EXCEL_FILE = "Unused_Keys_By-Days.xlsx"  # Workbook with account/user/key/reason.
RESULTS_CSV = "dry_run_securityservices_two_hop_results.csv"  # Output report.

SHARED_SERVICES_ACCOUNT_ID = "123912334908"  # SharedServices account.
SHARED_SERVICES_ROLE_NAME = "ManagedInstance-CrossAccountManagementRole"  # First role.
MEMBER_ACCOUNT_ROLE_NAME = "ManagedInstance-CrossAccountExecutionRole"  # Member role.
AWS_ACCOUNT_ID_LENGTH = 12  # AWS account IDs are always 12 digits.


# =========================
# PART 3: CLEAN ACCOUNT IDS
# =========================
def clean_account_id(value):
    account = str(value).strip()  # Convert Excel value to text and trim spaces.
    if account.endswith(".0"):  # Excel may format account IDs like 27063994883.0.
        account = account[:-2]  # Remove the trailing .0.
    account = "".join(char for char in account if char.isdigit())  # Keep digits only.
    return account.zfill(AWS_ACCOUNT_ID_LENGTH)  # Restore missing leading zeroes.


# =========================
# PART 4: LOAD WORKBOOK ROWS
# =========================
def load_rows():
    wb = load_workbook(EXCEL_FILE, data_only=True)  # Open workbook values.
    rows = []  # Store cleaned rows from every sheet.

    for sheet in wb.worksheets:  # Read each worksheet.
        headers = {}  # Map header name to column number.
        for col_num, cell in enumerate(sheet[1], start=1):  # Read row 1 headers.
            headers[str(cell.value).strip()] = col_num  # Example: account -> 1.

        for row_num in range(2, sheet.max_row + 1):  # Skip header row.
            account = sheet.cell(row=row_num, column=headers["account"]).value  # Account ID.
            user_name = sheet.cell(row=row_num, column=headers["user_name"]).value  # IAM user.
            access_key_id = sheet.cell(row=row_num, column=headers["access_key_id"]).value  # Key ID.
            reason = sheet.cell(row=row_num, column=headers["reason"]).value  # Reason.

            if not account or not user_name or not access_key_id:  # Ignore blanks.
                continue  # Move to the next row.

            rows.append({
                "sheet": sheet.title,  # Source sheet name.
                "row_number": row_num,  # Source Excel row.
                "account": clean_account_id(account),  # 12-digit AWS account.
                "user_name": str(user_name).strip(),  # Clean IAM user.
                "access_key_id": str(access_key_id).strip(),  # Clean key ID.
                "reason": str(reason).strip() if reason else "",  # Clean reason.
            })

    return rows  # Return all usable rows.


# =========================
# PART 5: AWS ASSUME ROLE HELPERS
# =========================
def assume_role(sts_client, role_arn, session_name):
    response = sts_client.assume_role(  # Ask STS for temporary role credentials.
        RoleArn=role_arn,  # Full ARN of the role to become.
        RoleSessionName=session_name,  # Name visible in CloudTrail.
    )
    return response["Credentials"]  # Return temporary access key/secret/token.


def make_client(service_name, credentials):
    return boto3.client(  # Build an AWS client from temporary credentials.
        service_name,  # Example: sts or iam.
        aws_access_key_id=credentials["AccessKeyId"],  # Temporary access key.
        aws_secret_access_key=credentials["SecretAccessKey"],  # Temporary secret.
        aws_session_token=credentials["SessionToken"],  # Temporary token.
    )


def assume_shared_services():
    sts = boto3.client("sts")  # Uses SecurityServices credentials from terminal.
    identity = sts.get_caller_identity()  # Confirm who is running the script.
    print(f"Starting identity: {identity['Arn']}")  # Show starting identity.

    role_arn = f"arn:aws:iam::{SHARED_SERVICES_ACCOUNT_ID}:role/{SHARED_SERVICES_ROLE_NAME}"
    print(f"Assuming SharedServices role: {role_arn}")  # Show first hop.
    return assume_role(sts, role_arn, "dry-run-sharedservices")  # SecurityServices -> SharedServices.


def assume_member_account(shared_credentials, account_id):
    shared_sts = make_client("sts", shared_credentials)  # STS client as SharedServices.
    role_arn = f"arn:aws:iam::{account_id}:role/{MEMBER_ACCOUNT_ROLE_NAME}"  # Member role ARN.
    print(f"  Assuming member account role: {role_arn}")  # Show second hop.
    member_credentials = assume_role(shared_sts, role_arn, f"dry-run-unused-keys-{account_id}")
    return make_client("iam", member_credentials)  # IAM client inside member account.


# =========================
# PART 6: DRY-RUN KEY CHECK
# =========================
def get_key_status(iam, user_name, access_key_id):
    response = iam.list_access_keys(UserName=user_name)  # Read keys for this IAM user.
    for key in response["AccessKeyMetadata"]:  # Check each key on the user.
        if key["AccessKeyId"] == access_key_id:  # Match workbook key.
            return key["Status"]  # Active or Inactive.
    return "NOT_FOUND"  # Workbook key was not found on that user.


# =========================
# PART 7: GROUP ROWS
# =========================
def group_rows_by_account(rows):
    grouped = defaultdict(list)  # account_id -> rows.
    for row in rows:  # Look at each workbook row.
        grouped[row["account"]].append(row)  # Put row under its account.
    return grouped  # Return grouped rows.


# =========================
# PART 8: MAIN PROGRAM
# =========================
def main():
    rows = load_rows()  # Read workbook.
    grouped_rows = group_rows_by_account(rows)  # Assume each account once.
    shared_credentials = assume_shared_services()  # First hop.
    results = []  # Store CSV rows.

    for account_id, account_rows in grouped_rows.items():  # Process account by account.
        print(f"\nChecking account {account_id}")  # Progress message.
        try:
            iam = assume_member_account(shared_credentials, account_id)  # Second hop.
            for row in account_rows:  # Check every key in this account.
                try:
                    key_status = get_key_status(iam, row["user_name"], row["access_key_id"])
                    if key_status == "Active":
                        action, status, error = "DRY_RUN_WOULD_DISABLE", "SUCCESS", ""
                    elif key_status == "Inactive":
                        action, status, error = "SKIP_ALREADY_INACTIVE", "SUCCESS", ""
                    else:
                        action, status, error = "SKIP_KEY_NOT_FOUND", "FAILED", "Access key not found for this user."
                except Exception as e:
                    key_status, action, status, error = "", "FAILED", "FAILED", str(e)

                results.append({**row, "key_status": key_status, "action": action, "status": status, "error": error, "timestamp": datetime.utcnow().isoformat()})
                print(f"{action}: {row['account']} {row['user_name']} {row['access_key_id']}")
        except Exception as e:
            for row in account_rows:
                results.append({**row, "key_status": "", "action": "FAILED_ASSUME_ROLE", "status": "FAILED", "error": str(e), "timestamp": datetime.utcnow().isoformat()})
            print(f"FAILED ACCOUNT {account_id}: {e}")

    with open(RESULTS_CSV, "w", newline="", encoding="utf-8") as f:  # Write report.
        writer = csv.DictWriter(f, fieldnames=list(results[0].keys()))  # Use result keys.
        writer.writeheader()  # Header row.
        writer.writerows(results)  # Data rows.

    print(f"\nChecked {len(results)}")  # Total checked.
    print(f"Results to {RESULTS_CSV}")  # Output file.


if __name__ == "__main__":  # Only runs when called as a script.
    main()  # Start program.
